from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_DIR = Path(__file__).resolve().parent.parent

INPUT_FILE = PROJECT_DIR / "Промежуточная таблица 2.xlsx"
OUTPUT_FILE = PROJECT_DIR / "Промежуточная таблица 3.xlsx"

DATA_SHEET_NAME = "Данные"
RESULT_SHEET_NAME = "3_Результат"
FINAL_SHEET_NAME = "4_Итоговые данные"
DO_RECALCULATED_SHEET = "5_ДО пересчитанное"
COUNTERPARTY_RECALCULATED_SHEET = "6_Контрагент пересчитанное"


def create_intermediate_table_3():
    workbook = load_workbook(INPUT_FILE)
    data_sheet = workbook[DATA_SHEET_NAME]

    # Цвета
    title_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    rows = data_sheet.iter_rows(values_only=True)
    headers = list(next(rows))

    records = []

    for row in rows:
        if all(value is None for value in row):
            continue

        record = dict(zip(headers, row))

        record["Остаток ДЗ"] = record["Сумма ДЗ"] or 0
        record["Остаток КЗ"] = record["Сумма КЗ"] or 0

        records.append(record)

    grouped_by_counterparty = defaultdict(list)

    for record in records:
        grouped_by_counterparty[
            record["Контрагент"]
        ].append(record)

    operations = []
    counterparties_without_offset = []
    operation_number = 1

    # Расчёт взаимозачётов
    for counterparty, counterparty_records in grouped_by_counterparty.items():
        receivables = [
            record
            for record in counterparty_records
            if record["Остаток ДЗ"] > 0
        ]

        payables = [
            record
            for record in counterparty_records
            if record["Остаток КЗ"] > 0
        ]

        offset_found = False

        for payable in payables:
            for receivable in receivables:
                if payable["Остаток КЗ"] <= 0:
                    break

                if receivable["Остаток ДЗ"] <= 0:
                    continue

                kz_before = payable["Остаток КЗ"]
                dz_before = receivable["Остаток ДЗ"]

                offset_amount = min(
                    kz_before,
                    dz_before,
                )

                payable["Остаток КЗ"] -= offset_amount
                receivable["Остаток ДЗ"] -= offset_amount

                kz_after = payable["Остаток КЗ"]
                dz_after = receivable["Остаток ДЗ"]

                if payable["ДО"] == receivable["ДО"]:
                    offset_type = "Двусторонний"
                else:
                    offset_type = "Трёхсторонний"

                movement = (
                    f'{payable["ДО"]} → '
                    f'{counterparty} → '
                    f'{receivable["ДО"]}'
                )

                operations.append({
                    "Номер": operation_number,
                    "Контрагент": counterparty,
                    "Вид зачёта": offset_type,
                    "Схема": movement,
                    "Сумма зачёта": offset_amount,

                    "ID КЗ": payable["ID"],
                    "ДО КЗ": payable["ДО"],
                    "Договор КЗ": payable["Договор"],
                    "КЗ до": kz_before,
                    "КЗ после": kz_after,

                    "ID ДЗ": receivable["ID"],
                    "ДО ДЗ": receivable["ДО"],
                    "Договор ДЗ": receivable["Договор"],
                    "Вид ДЗ": receivable["Вид ДЗ"],
                    "ДЗ до": dz_before,
                    "ДЗ после": dz_after,
                })

                operation_number += 1
                offset_found = True

        if not offset_found:
            total_receivable = sum(
                record["Сумма ДЗ"] or 0
                for record in counterparty_records
            )

            total_payable = sum(
                record["Сумма КЗ"] or 0
                for record in counterparty_records
            )

            if total_receivable == 0:
                reason = "Нет ДЗ"
            elif total_payable == 0:
                reason = "Нет КЗ"
            else:
                reason = "Зачёт не найден"

            counterparties_without_offset.append([
                counterparty,
                total_receivable,
                total_payable,
                reason,
            ])

    # ---------------------------------------------------------
    # Лист 3_Результат
    # ---------------------------------------------------------

    if RESULT_SHEET_NAME in workbook.sheetnames:
        del workbook[RESULT_SHEET_NAME]

    result_sheet = workbook.create_sheet(
        RESULT_SHEET_NAME,
        index=3,
    )

    result_sheet.append([
        "Возможные взаимозачёты"
    ])

    result_sheet.append([
        "№",
        "Строка",
        "Контрагент",
        "Вид",
        "Движение",
        "Договор",
        "Сумма / изменение",
        "ID источника",
    ])

    for operation in operations:
        result_sheet.append([
            operation["Номер"],
            "ОПЕРАЦИЯ ВЗАИМОЗАЧЁТА",
            operation["Контрагент"],
            operation["Вид зачёта"],
            operation["Схема"],
            "",
            operation["Сумма зачёта"],
            "",
        ])

        result_sheet.append([
            "",
            "КЗ",
            operation["Контрагент"],
            "Кредиторская задолженность",
            (
                f'{operation["ДО КЗ"]} должна '
                f'{operation["Контрагент"]}'
            ),
            operation["Договор КЗ"],
            (
                f'{operation["КЗ до"]} → '
                f'{operation["КЗ после"]}'
            ),
            operation["ID КЗ"],
        ])

        result_sheet.append([
            "",
            "ДЗ",
            operation["Контрагент"],
            (
                "Дебиторская задолженность "
                f'({operation["Вид ДЗ"]})'
            ),
            (
                f'{operation["Контрагент"]} должна '
                f'{operation["ДО ДЗ"]}'
            ),
            operation["Договор ДЗ"],
            (
                f'{operation["ДЗ до"]} → '
                f'{operation["ДЗ после"]}'
            ),
            operation["ID ДЗ"],
        ])

        result_sheet.append([None])

    result_sheet.append([
        "Контрагенты без возможного зачёта"
    ])

    result_sheet.append([
        "Контрагент",
        "Общая ДЗ",
        "Общая КЗ",
        "Причина",
    ])

    for counterparty in counterparties_without_offset:
        result_sheet.append(counterparty)

    # ---------------------------------------------------------
    # Лист 4_Итоговые данные
    # ---------------------------------------------------------

    if FINAL_SHEET_NAME in workbook.sheetnames:
        del workbook[FINAL_SHEET_NAME]

    final_sheet = workbook.create_sheet(
        FINAL_SHEET_NAME,
        index=4,
    )

    final_sheet.append(headers)
    final_sheet.freeze_panes = "A2"

    for record in records:
        final_row = []

        for header in headers:
            if header == "Сумма ДЗ":
                final_row.append(
                    record["Остаток ДЗ"]
                )
            elif header == "Сумма КЗ":
                final_row.append(
                    record["Остаток КЗ"]
                )
            else:
                final_row.append(
                    record[header]
                )

        final_sheet.append(final_row)

    final_sheet.auto_filter.ref = (
        final_sheet.dimensions
    )

    # ---------------------------------------------------------
    # Лист 5_ДО пересчитанное
    # ---------------------------------------------------------

    if DO_RECALCULATED_SHEET in workbook.sheetnames:
        del workbook[DO_RECALCULATED_SHEET]

    do_sheet = workbook.create_sheet(
        DO_RECALCULATED_SHEET,
        index=5,
    )

    grouped_by_do = defaultdict(list)

    for record in records:
        grouped_by_do[
            record["ДО"]
        ].append(record)

    do_headers = [
        "ID",
        "Контрагент",
        "Договор",
        "Вид ДЗ",
        "Сумма ДЗ до зачёта",
        "Сумма КЗ до зачёта",
        "Остаток ДЗ после зачёта",
        "Остаток КЗ после зачёта",
        "Кто кому должен после зачёта",
    ]

    for subsidiary, subsidiary_records in grouped_by_do.items():
        # Голубая строка с названием ДО
        do_sheet.append([
            f"ДО: {subsidiary}"
        ])

        title_row = do_sheet.max_row

        # Заливка без объединения ячеек
        for column_number in range(
            1,
            len(do_headers) + 1,
        ):
            cell = do_sheet.cell(
                row=title_row,
                column=column_number,
            )

            cell.fill = title_fill

        do_sheet.cell(
            row=title_row,
            column=1,
        ).font = Font(
            bold=True,
            size=12,
        )

        # Зелёная строка заголовков
        do_sheet.append(do_headers)

        header_row = do_sheet.max_row

        for cell in do_sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        total_receivable_before = 0
        total_payable_before = 0
        total_receivable_after = 0
        total_payable_after = 0

        for record in subsidiary_records:
            counterparty = record["Контрагент"]

            receivable_before = (
                record["Сумма ДЗ"] or 0
            )

            payable_before = (
                record["Сумма КЗ"] or 0
            )

            receivable_after = (
                record["Остаток ДЗ"]
            )

            payable_after = (
                record["Остаток КЗ"]
            )

            if receivable_after > 0 and payable_after > 0:
                debt_direction = (
                    f"{counterparty} должен {subsidiary}: "
                    f"{receivable_after}; "
                    f"{subsidiary} должно {counterparty}: "
                    f"{payable_after}"
                )
            elif receivable_after > 0:
                debt_direction = (
                    f"{counterparty} должен "
                    f"{subsidiary}: {receivable_after}"
                )
            elif payable_after > 0:
                debt_direction = (
                    f"{subsidiary} должно "
                    f"{counterparty}: {payable_after}"
                )
            else:
                debt_direction = "Задолженность погашена"

            do_sheet.append([
                record["ID"],
                counterparty,
                record["Договор"],
                record["Вид ДЗ"],
                receivable_before,
                payable_before,
                receivable_after,
                payable_after,
                debt_direction,
            ])

            total_receivable_before += (
                receivable_before
            )

            total_payable_before += (
                payable_before
            )

            total_receivable_after += (
                receivable_after
            )

            total_payable_after += (
                payable_after
            )

        # Жёлтая итоговая строка
        do_sheet.append([
            "Итого по ДО",
            "",
            "",
            "",
            total_receivable_before,
            total_payable_before,
            total_receivable_after,
            total_payable_after,
            "",
        ])

        total_row = do_sheet.max_row

        for cell in do_sheet[total_row]:
            cell.font = Font(bold=True)
            cell.fill = total_fill

        do_sheet.append([None])

    # ---------------------------------------------------------
    # Лист 6_Контрагент пересчитанное
    # ---------------------------------------------------------

    if COUNTERPARTY_RECALCULATED_SHEET in workbook.sheetnames:
        del workbook[COUNTERPARTY_RECALCULATED_SHEET]

    counterparty_sheet = workbook.create_sheet(
        COUNTERPARTY_RECALCULATED_SHEET,
        index=6,
    )

    counterparty_headers = [
        "ID",
        "ДО",
        "Договор",
        "Вид ДЗ",
        "Сумма ДЗ до зачёта",
        "Сумма КЗ до зачёта",
        "Остаток ДЗ после зачёта",
        "Остаток КЗ после зачёта",
    ]

    for counterparty, counterparty_records in grouped_by_counterparty.items():
        # Голубая строка с названием контрагента
        counterparty_sheet.append([
            f"Контрагент: {counterparty}"
        ])

        title_row = counterparty_sheet.max_row

        # Заливка без объединения ячеек
        for column_number in range(
            1,
            len(counterparty_headers) + 1,
        ):
            cell = counterparty_sheet.cell(
                row=title_row,
                column=column_number,
            )

            cell.fill = title_fill

        counterparty_sheet.cell(
            row=title_row,
            column=1,
        ).font = Font(
            bold=True,
            size=12,
        )

        # Зелёная строка заголовков
        counterparty_sheet.append(
            counterparty_headers
        )

        header_row = counterparty_sheet.max_row

        for cell in counterparty_sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        total_receivable_before = 0
        total_payable_before = 0
        total_receivable_after = 0
        total_payable_after = 0

        for record in counterparty_records:
            receivable_before = (
                record["Сумма ДЗ"] or 0
            )

            payable_before = (
                record["Сумма КЗ"] or 0
            )

            receivable_after = (
                record["Остаток ДЗ"]
            )

            payable_after = (
                record["Остаток КЗ"]
            )

            counterparty_sheet.append([
                record["ID"],
                record["ДО"],
                record["Договор"],
                record["Вид ДЗ"],
                receivable_before,
                payable_before,
                receivable_after,
                payable_after,
            ])

            total_receivable_before += (
                receivable_before
            )

            total_payable_before += (
                payable_before
            )

            total_receivable_after += (
                receivable_after
            )

            total_payable_after += (
                payable_after
            )

        # Жёлтая итоговая строка
        counterparty_sheet.append([
            "Итого",
            "",
            "",
            "",
            total_receivable_before,
            total_payable_before,
            total_receivable_after,
            total_payable_after,
        ])

        total_row = counterparty_sheet.max_row

        for cell in counterparty_sheet[total_row]:
            cell.font = Font(bold=True)
            cell.fill = total_fill

        counterparty_sheet.append([None])

    # ---------------------------------------------------------
    # Автоматическая ширина столбцов
    # ---------------------------------------------------------

    sheets_for_auto_width = [
        result_sheet,
        final_sheet,
        do_sheet,
        counterparty_sheet,
    ]

    for worksheet in sheets_for_auto_width:
        for column_cells in worksheet.iter_cols():
            max_length = 0

            for cell in column_cells:
                if cell.value is not None:
                    cell_length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        cell_length,
                    )

            column_letter = get_column_letter(
                column_cells[0].column
            )

            worksheet.column_dimensions[
                column_letter
            ].width = max_length + 2

    workbook.save(OUTPUT_FILE)
    workbook.close()

    print(f"Создан файл: {OUTPUT_FILE}")


if __name__ == "__main__":
    create_intermediate_table_3()