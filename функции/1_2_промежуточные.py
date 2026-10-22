from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


PROJECT_DIR = Path(__file__).resolve().parent.parent

INPUT_FILE = PROJECT_DIR / "задолженность.xlsx"

OUTPUT_FILE_1 = (
    PROJECT_DIR
    / "Промежуточная таблица 1.xlsx"
)

OUTPUT_FILE_2 = (
    PROJECT_DIR
    / "Промежуточная таблица 2.xlsx"
)

DATA_SHEET_NAME = "Данные"
DO_SHEET_NAME = "1_ДО"
COUNTERPARTY_SHEET_NAME = "2_Контрагент"

REQUIRED_COLUMNS = [
    "ID",
    "ДО",
    "Контрагент",
    "Договор",
    "Вид ДЗ",
    "Сумма ДЗ",
    "Сумма КЗ",
]


def read_excel_data(workbook):
    """
    Читает первичные данные с листа «Данные».
    """

    if DATA_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Не найден лист «{DATA_SHEET_NAME}»."
        )

    worksheet = workbook[DATA_SHEET_NAME]
    rows = worksheet.iter_rows(values_only=True)

    headers = next(rows, None)

    if headers is None:
        raise ValueError(
            "Лист с первичными данными пустой."
        )

    headers = list(headers)

    if headers != REQUIRED_COLUMNS:
        raise ValueError(
            "Неверные столбцы.\n"
            f"Ожидается: {REQUIRED_COLUMNS}\n"
            f"Получено: {headers}"
        )

    records = []

    for values in rows:
        if all(value is None for value in values):
            continue

        record = dict(
            zip(REQUIRED_COLUMNS, values)
        )

        records.append(record)

    return records


def get_debt_direction(record):
    """
    Возвращает текст: кто кому должен.
    """

    subsidiary = record["ДО"]
    counterparty = record["Контрагент"]

    receivable = record["Сумма ДЗ"] or 0
    payable = record["Сумма КЗ"] or 0

    if receivable > 0 and payable > 0:
        return (
            f"{counterparty} должен {subsidiary}: "
            f"{receivable}; "
            f"{subsidiary} должно {counterparty}: "
            f"{payable}"
        )

    if receivable > 0:
        return (
            f"{counterparty} должен "
            f"{subsidiary}: {receivable}"
        )

    if payable > 0:
        return (
            f"{subsidiary} должно "
            f"{counterparty}: {payable}"
        )

    return "Задолженность отсутствует"


def add_do_sheet(workbook, records):
    """
    Добавляет лист «1_ДО».
    """

    if DO_SHEET_NAME in workbook.sheetnames:
        del workbook[DO_SHEET_NAME]

    worksheet = workbook.create_sheet(
        DO_SHEET_NAME
    )

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    grouped_records = defaultdict(list)

    for record in records:
        grouped_records[
            record["ДО"]
        ].append(record)

    headers = [
        "ID",
        "Контрагент",
        "Договор",
        "Вид ДЗ",
        "Сумма ДЗ (контрагент должен ДО)",
        "Сумма КЗ (ДО должно контрагенту)",
        "Кто кому должен",
    ]

    current_row = 1

    for subsidiary, subsidiary_records in grouped_records.items():
        # Название ДО
        worksheet.cell(
            row=current_row,
            column=1,
            value=f"ДО: {subsidiary}",
        )

        # Голубая строка без объединения ячеек
        for column_number in range(
            1,
            len(headers) + 1,
        ):
            cell = worksheet.cell(
                row=current_row,
                column=column_number,
            )

            cell.fill = title_fill

        worksheet.cell(
            row=current_row,
            column=1,
        ).font = Font(
            bold=True,
            size=12,
        )

        current_row += 1

        # Заголовки
        for column_number, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=current_row,
                column=column_number,
                value=header,
            )

            cell.font = Font(bold=True)
            cell.fill = header_fill

        current_row += 1

        total_receivable = 0
        total_payable = 0

        for record in subsidiary_records:
            receivable = (
                record["Сумма ДЗ"] or 0
            )

            payable = (
                record["Сумма КЗ"] or 0
            )

            values = [
                record["ID"],
                record["Контрагент"],
                record["Договор"],
                record["Вид ДЗ"],
                receivable,
                payable,
                get_debt_direction(record),
            ]

            for column_number, value in enumerate(
                values,
                start=1,
            ):
                worksheet.cell(
                    row=current_row,
                    column=column_number,
                    value=value,
                )

            total_receivable += receivable
            total_payable += payable

            current_row += 1

        # Итоговая строка
        worksheet.cell(
            row=current_row,
            column=1,
            value="Итого по ДО",
        )

        worksheet.cell(
            row=current_row,
            column=5,
            value=total_receivable,
        )

        worksheet.cell(
            row=current_row,
            column=6,
            value=total_payable,
        )

        for column_number in range(
            1,
            len(headers) + 1,
        ):
            cell = worksheet.cell(
                row=current_row,
                column=column_number,
            )

            cell.font = Font(bold=True)
            cell.fill = total_fill

        current_row += 2

    column_widths = {
        "A": 10,
        "B": 20,
        "C": 15,
        "D": 12,
        "E": 34,
        "F": 34,
        "G": 55,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[
            column
        ].width = width


def add_counterparty_sheet(workbook, records):
    """
    Добавляет лист «2_Контрагент».
    """

    if COUNTERPARTY_SHEET_NAME in workbook.sheetnames:
        del workbook[COUNTERPARTY_SHEET_NAME]

    worksheet = workbook.create_sheet(
        COUNTERPARTY_SHEET_NAME
    )

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    grouped_records = defaultdict(list)

    for record in records:
        grouped_records[
            record["Контрагент"]
        ].append(record)

    headers = [
        "ID",
        "ДО",
        "Договор",
        "Вид ДЗ",
        "Сумма ДЗ",
        "Сумма КЗ",
        "Остаток ДЗ",
        "Остаток КЗ",
    ]

    for counterparty, counterparty_records in grouped_records.items():
        # Название контрагента
        worksheet.append([
            f"Контрагент: {counterparty}"
        ])

        title_row = worksheet.max_row

        # Голубая строка без объединения ячеек
        for column_number in range(
            1,
            len(headers) + 1,
        ):
            cell = worksheet.cell(
                row=title_row,
                column=column_number,
            )

            cell.fill = title_fill

        worksheet.cell(
            row=title_row,
            column=1,
        ).font = Font(
            bold=True,
            size=12,
        )

        # Заголовки
        worksheet.append(headers)

        header_row = worksheet.max_row

        for cell in worksheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        total_receivable = 0
        total_payable = 0

        for record in counterparty_records:
            receivable = (
                record["Сумма ДЗ"] or 0
            )

            payable = (
                record["Сумма КЗ"] or 0
            )

            worksheet.append([
                record["ID"],
                record["ДО"],
                record["Договор"],
                record["Вид ДЗ"],
                receivable,
                payable,
                receivable,
                payable,
            ])

            total_receivable += receivable
            total_payable += payable

        # Итоговая строка
        worksheet.append([
            "Итого",
            "",
            "",
            "",
            total_receivable,
            total_payable,
            total_receivable,
            total_payable,
        ])

        total_row = worksheet.max_row

        for cell in worksheet[total_row]:
            cell.font = Font(bold=True)
            cell.fill = total_fill

        worksheet.append([None])

    column_widths = {
        "A": 12,
        "B": 12,
        "C": 15,
        "D": 12,
        "E": 16,
        "F": 16,
        "G": 18,
        "H": 18,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[
            column
        ].width = width


def create_intermediate_files():
    """
    Последовательно создаёт первую и вторую
    промежуточные книги.
    """

    if not INPUT_FILE.exists():
        print(f"Не найден файл: {INPUT_FILE}")
        return

    workbook = load_workbook(INPUT_FILE)

    try:
        records = read_excel_data(workbook)

        # Создаём лист 1_ДО
        add_do_sheet(
            workbook,
            records,
        )

        # Сохраняем первую промежуточную книгу
        workbook.save(OUTPUT_FILE_1)

        print(
            f"Создан файл: {OUTPUT_FILE_1}"
        )

        # Добавляем лист 2_Контрагент
        add_counterparty_sheet(
            workbook,
            records,
        )

        # Сохраняем вторую промежуточную книгу
        workbook.save(OUTPUT_FILE_2)

        print(
            f"Создан файл: {OUTPUT_FILE_2}"
        )

    finally:
        workbook.close()


if __name__ == "__main__":
    create_intermediate_files()