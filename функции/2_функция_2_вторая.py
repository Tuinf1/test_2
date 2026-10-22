from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


PROJECT_DIR = Path(__file__).resolve().parent.parent

INPUT_FILE = PROJECT_DIR / "Промежуточная таблица 1.xlsx"
OUTPUT_FILE = PROJECT_DIR / "Промежуточная таблица 2.xlsx"

DATA_SHEET = "Данные"
NEW_SHEET = "2_Контрагент"


def create_intermediate_table_2():
    workbook = load_workbook(INPUT_FILE)

    # Читаем первичные данные
    data_sheet = workbook[DATA_SHEET]
    rows = data_sheet.iter_rows(values_only=True)

    headers = next(rows)

    records = [
        dict(zip(headers, row))
        for row in rows
        if any(value is not None for value in row)
    ]

    # Удаляем старую версию листа
    if NEW_SHEET in workbook.sheetnames:
        del workbook[NEW_SHEET]

    worksheet = workbook.create_sheet(
        NEW_SHEET
    )

    # Цвета
    title_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    # Группируем строки по контрагентам
    grouped_records = defaultdict(list)

    for record in records:
        grouped_records[
            record["Контрагент"]
        ].append(record)

    table_headers = [
        "ID",
        "ДО",
        "Договор",
        "Вид ДЗ",
        "Сумма ДЗ",
        "Сумма КЗ",
        "Остаток ДЗ",
        "Остаток КЗ",
    ]

    for counterparty, counterparty_records in grouped_records.items():
        # Строка с названием контрагента
        worksheet.append([
            f"Контрагент: {counterparty}"
        ])

        title_row = worksheet.max_row

        # Красим строку без объединения ячеек
        for column_number in range(
            1,
            len(table_headers) + 1,
        ):
            cell = worksheet.cell(
                row=title_row,
                column=column_number,
            )

            cell.fill = title_fill

        # Название контрагента находится только в первой ячейке
        title_cell = worksheet.cell(
            row=title_row,
            column=1,
        )

        title_cell.font = Font(
            bold=True,
            size=12,
        )

        # Заголовки таблицы
        worksheet.append(table_headers)

        header_row = worksheet.max_row

        for cell in worksheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        total_receivable = 0
        total_payable = 0

        for record in counterparty_records:
            receivable = (
                record["Сумма ДЗ"] or 0
            )

            payable = (
                record["Сумма КЗ"] or 0
            )

            worksheet.append([
                record["ID"],
                record["ДО"],
                record["Договор"],
                record["Вид ДЗ"],
                receivable,
                payable,
                receivable,
                payable,
            ])

            total_receivable += receivable
            total_payable += payable

        # Итоговая строка
        worksheet.append([
            "Итого",
            "",
            "",
            "",
            total_receivable,
            total_payable,
            total_receivable,
            total_payable,
        ])

        total_row = worksheet.max_row

        for cell in worksheet[total_row]:
            cell.font = Font(bold=True)
            cell.fill = total_fill

        # Пустая строка между контрагентами
        worksheet.append([None])

    # Ширина столбцов
    column_widths = {
        "A": 12,
        "B": 12,
        "C": 15,
        "D": 12,
        "E": 16,
        "F": 16,
        "G": 18,
        "H": 18,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    workbook.save(OUTPUT_FILE)
    workbook.close()

    print(f"Создан файл: {OUTPUT_FILE}")


if __name__ == "__main__":
    create_intermediate_table_2()