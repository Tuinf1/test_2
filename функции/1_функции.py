# чтение и проверка данных

from pathlib import Path
from openpyxl import load_workbook


REQUIRED_COLUMNS = [
    "ID",
    "ДО",
    "Контрагент",
    "Договор",
    "Вид ДЗ",
    "Сумма ДЗ",
    "Сумма КЗ",
]


def is_number(value) -> bool:
    """Проверяет, что значение является числом."""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def validate_headers(headers: list) -> None:
    """Проверяет, что Excel содержит нужные столбцы."""
    if headers != REQUIRED_COLUMNS:
        raise ValueError(
            "Неверная структура столбцов.\n"
            f"Ожидается: {REQUIRED_COLUMNS}\n"
            f"Получено: {headers}"
        )


def validate_record(record: dict, row_number: int, used_ids: set) -> list[str]:
    """Проверяет одну строку входных данных."""
    errors = []

    # Проверка обязательных полей
    for field in ["ID", "ДО", "Контрагент", "Договор", "Вид ДЗ"]:
        if record[field] in (None, ""):
            errors.append(
                f"Строка {row_number}: не заполнено поле «{field}»."
            )

    # Проверка ID
    record_id = record["ID"]

    if not isinstance(record_id, int) or record_id <= 0:
        errors.append(
            f"Строка {row_number}: ID должен быть положительным целым числом."
        )
    elif record_id in used_ids:
        errors.append(
            f"Строка {row_number}: найден дублирующийся ID {record_id}."
        )
    else:
        used_ids.add(record_id)

    # Проверка вида дебиторской задолженности
    if record["Вид ДЗ"] not in ("ПДЗ", "ДЗ"):
        errors.append(
            f"Строка {row_number}: «Вид ДЗ» должен быть «ПДЗ» или «ДЗ»."
        )

    # Проверка сумм
    for field in ["Сумма ДЗ", "Сумма КЗ"]:
        value = record[field]

        if not is_number(value):
            errors.append(
                f"Строка {row_number}: «{field}» должна быть числом."
            )
        elif value < 0:
            errors.append(
                f"Строка {row_number}: «{field}» не может быть отрицательной."
            )

    return errors


def read_and_validate_excel(input_path: str | Path) -> list[dict]:
    """
    Читает файл Excel и возвращает список записей.

    Ожидаемый файл:
    - путь: input/задолженность.xlsx;
    - лист: Данные;
    - столбцы: ID, ДО, Контрагент, Договор,
      Вид ДЗ, Сумма ДЗ, Сумма КЗ.
    """
    file_path = Path(input_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Входной файл не найден: {file_path}"
        )

    workbook = load_workbook(file_path, data_only=True)

    try:
        if "Данные" not in workbook.sheetnames:
            raise ValueError("В файле отсутствует лист «Данные».")

        worksheet = workbook["Данные"]

        # Чтение и проверка заголовков
        headers = [cell.value for cell in worksheet[1]]
        validate_headers(headers)

        records = []
        errors = []
        used_ids = set()

        # Чтение строк, начиная со второй
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            # Полностью пустые строки игнорируем
            if all(value is None for value in values):
                continue

            record = dict(zip(REQUIRED_COLUMNS, values))

            # Удаляем случайные пробелы в текстовых полях
            for field in ["ДО", "Контрагент", "Договор", "Вид ДЗ"]:
                if isinstance(record[field], str):
                    record[field] = record[field].strip()

            errors.extend(
                validate_record(record, row_number, used_ids)
            )

            records.append(record)

        if not records:
            errors.append("В файле нет строк с данными.")

        if errors:
            raise ValueError(
                "Ошибки во входном файле:\n- "
                + "\n- ".join(errors)
            )

        return records

    finally:
        workbook.close()



if __name__ == "__main__":
    test_file = "input/задолженность.xlsx"

    try:
        records = read_and_validate_excel(test_file)

        print("Проверка пройдена успешно.")
        print(f"Количество строк: {len(records)}")
        print("Первые 3 записи:")

        for record in records[:3]:
            print(record)

    except (FileNotFoundError, ValueError) as error:
        print(f"Ошибка проверки: {error}")